from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "manual-probiolab"
CAP = OUT / "capturas"
MANUAL = OUT / "Manual_de_Usuario_Bio_Probiolab_v1.0.docx"
REPORT = OUT / "Reporte_de_Pruebas_y_Correcciones_Bio_Probiolab_v1.0.docx"
TODAY = "8 de septiembre de 2026"
NAVY = "17465B"
TEAL = "1E7890"
PALE = "EAF4F7"
GRID = "D9D9D9"
INK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_repeat_table_header(table):
    repeat_header(table.rows[0])


def set_picture_alt(inline_shape, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def prevent_table_row_splits(doc):
    """Keep each logical table record on one page when Word paginates."""
    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:cantSplit")) is None:
                cant_split = OxmlElement("w:cantSplit")
                cant_split.set(qn("w:val"), "true")
                tr_pr.append(cant_split)


def add_field(paragraph, instruction, fallback=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = RGBColor(32, 42, 46)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    for name, size, before, after in (("Title", 28, 0, 16), ("Heading 1", 20, 18, 8), ("Heading 2", 15, 14, 6), ("Heading 3", 12, 10, 4)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(8.5)
    styles["Caption"].font.italic = True
    styles["Caption"].font.color.rgb = RGBColor(70, 80, 84)
    styles["Caption"].paragraph_format.space_before = Pt(3)
    styles["Caption"].paragraph_format.space_after = Pt(9)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_header_footer(doc, short_title):
    for section in doc.sections:
        header = section.header
        p = header.paragraphs[0]
        p.text = short_title
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in p.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = INK
        footer = section.footer
        table = footer.add_table(rows=1, cols=3, width=Inches(7.0))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [2.7, 1.6, 2.7]
        for idx, width in enumerate(widths):
            table.cell(0, idx).width = Inches(width)
            set_cell_margins(table.cell(0, idx), 0, 40, 0, 40)
        left = table.cell(0, 0).paragraphs[0]
        left.add_run("PROBIOLAB  |  Uso interno").font.size = Pt(8)
        middle = table.cell(0, 1).paragraphs[0]
        middle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        middle.add_run("v1.0").font.size = Pt(8)
        right = table.cell(0, 2).paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right.add_run("Página ").font.size = Pt(8)
        add_field(right, "PAGE", "1")
        right.add_run(" de ").font.size = Pt(8)
        add_field(right, "NUMPAGES", "1")
        for cell in table.rows[0].cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "nil")
            borders.append(el)
        tbl_pr.append(borders)


def add_cover(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    logo = ROOT / "assets" / "probiolab-logo-print-transparent.png"
    if logo.exists():
        shape = p.add_run().add_picture(str(logo), width=Inches(3.0))
        set_picture_alt(shape, "Logotipo de PROBIOLAB")
    title_p = doc.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.add_run(title)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(subtitle).bold = True
    sub.runs[0].font.size = Pt(14)
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    labels = [("Sistema", "Bio / PROBIOLAB"), ("Versión del documento", "1.0"), ("Fecha", TODAY), ("Clasificación", "Uso interno operativo y administrativo")]
    for i, (label, value) in enumerate(labels):
        meta.cell(i, 0).text = label
        meta.cell(i, 1).text = value
        meta.cell(i, 0).width = Inches(2.0)
        meta.cell(i, 1).width = Inches(4.2)
        meta.cell(i, 0).paragraphs[0].runs[0].bold = True
        set_cell_shading(meta.cell(i, 0), PALE)
        for cell in meta.rows[i].cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(meta)
    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run("Documento validado contra la interfaz local de PROBIOLAB y escenarios de prueba aislados.")
    doc.add_page_break()


def add_toc(doc):
    doc.add_heading("Contenido", level=1)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Contenido automático")
    doc.add_paragraph("En Microsoft Word, seleccione la tabla y use Actualizar tabla si el documento fue modificado después de su emisión.")
    doc.add_page_break()


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    r.bold = True
    p.add_run(text)
    return p


def add_numbered_steps(doc, steps):
    for step in steps:
        p = doc.add_paragraph(style="List Number")
        p.add_run(step)


def add_figure(doc, filename, caption, number):
    path = CAP / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(6.85))
    set_picture_alt(shape, caption)
    c = doc.add_paragraph(style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.add_run(f"Figura {number}. {caption}")
    return number + 1


def add_procedure(doc, title, objective, roles, prerequisites, route, steps, expected, validations, errors, evidence, impact, figures, fig_no):
    doc.add_heading(title, level=2)
    add_label_paragraph(doc, "Objetivo", objective)
    add_label_paragraph(doc, "Quién puede realizarlo", roles)
    add_label_paragraph(doc, "Requisitos previos", prerequisites)
    add_label_paragraph(doc, "Ruta", route)
    doc.add_heading("Paso a paso", level=3)
    add_numbered_steps(doc, steps)
    for filename, caption in figures:
        fig_no = add_figure(doc, filename, caption, fig_no)
    add_label_paragraph(doc, "Resultado esperado", expected)
    add_label_paragraph(doc, "Validaciones y alertas", validations)
    add_label_paragraph(doc, "Errores comunes y solución", errors)
    add_label_paragraph(doc, "Evidencias y documentos", evidence)
    add_label_paragraph(doc, "Impacto posterior", impact)
    return fig_no


def add_flow_table(doc):
    stages = [
        ("1", "Cotización", "PDF y Excel", "OC del cliente"),
        ("2", "Aceptación u OC", "Referencia, archivo y huella", "Pedido activo"),
        ("3", "Pedido y requisición", "Partidas y cantidades", "Necesidades de compra"),
        ("4", "OC proveedor", "Asignaciones por pedido", "Recepción"),
        ("5", "Recepción", "Remisión, lote y evidencia", "Entrada y disponibilidad"),
        ("6", "Inventario y reserva", "Movimiento y reserva", "Surtido"),
        ("7", "Remisión y salida", "Folio, salida y evidencia", "Entrega"),
        ("8", "Factura y cierre", "Factura y expediente", "Cierre auditable"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["Etapa", "Proceso", "Evidencia conservada", "Habilita"]
    for j, value in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = value
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
    widths = [0.55, 1.75, 2.75, 1.55]
    for row_index, row in enumerate(stages, start=1):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = value
            cells[j].width = Inches(widths[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[j])
            if row_index % 2 == 0:
                set_cell_shading(cells[j], "F3F8FA")
    set_repeat_table_header(table)
    set_table_borders(table)


def build_manual():
    doc = Document()
    configure_styles(doc)
    add_cover(doc, "Manual de Usuario Bio Probiolab", "Flujo operativo desde cotización hasta entrega al cliente")
    add_toc(doc)
    fig = 1

    doc.add_heading("1 Introducción", level=1)
    doc.add_paragraph("PROBIOLAB controla cotizaciones, pedidos de clientes, compras a proveedores, almacenes, seguimiento, remisiones, documentos y auditoría. Este manual explica la versión validada el 8 de septiembre de 2026. Está dirigido a personal operativo, vendedores, responsables de compras y almacén, supervisores, administradores y auditores.")
    add_label_paragraph(doc, "Alcance", "Interfaz web local, datos de demostración, cuatro perfiles configurables y el flujo documental disponible en la versión 1.0. La conexión central PostgreSQL requiere configuración de infraestructura y no estuvo disponible en el equipo de validación.")
    add_label_paragraph(doc, "Campos obligatorios", "El asterisco indica que el campo debe completarse. Cuando falta información, el sistema marca el campo en rojo, agrega una explicación accesible y lleva al primer error.")
    add_label_paragraph(doc, "Evidencias", "Los archivos admitidos se guardan con nombre, tipo, fecha y huella cuando el flujo formal lo exige. Los documentos del área TEST no tienen validez operativa.")
    add_label_paragraph(doc, "Estados y fechas", "Las fechas compromiso vencidas aparecen como alertas. Los estados de pedido se derivan de cantidades compradas, recibidas, reservadas, entregadas, facturadas o canceladas.")
    fig = add_figure(doc, "01-resumen-general.png", "Resumen general con periodo, indicadores y accesos a la operación", fig)

    doc.add_heading("2 Flujo general de Bio", level=1)
    doc.add_paragraph("Cada etapa conserva la referencia que originó a la siguiente. El pedido de venta nace de una cotización; la compra conserva la asignación de cada partida; la recepción genera disponibilidad; la remisión produce la salida y la entrega queda vinculada al expediente.")
    add_flow_table(doc)
    fig = add_figure(doc, "11-resumen-seguimiento.png", "Resumen del avance por pedido y por partida", fig)

    doc.add_heading("3 Navegación acceso y áreas de trabajo", level=1)
    fig = add_procedure(doc, "3.1 Navegar por el sistema", "Ubicar un módulo, volver a su resumen y consultar alertas.", "Todos los perfiles con permiso de consulta.", "Abrir PROBIOLAB desde el servidor en http://127.0.0.1:8787. No abrir index.html directamente.", "Menú lateral o búsqueda global.", ["Use los grupos Cotizaciones, Almacén, Seguimiento y remisiones, y Productos.", "Seleccione el resumen del grupo o una opción secundaria.", "Use la búsqueda superior para localizar productos, movimientos, cotizaciones, pedidos o proveedores.", "Revise el indicador inferior: Base central conectada o API sin conexión."], "La vista seleccionada queda resaltada y se conserva al recargar.", "Las opciones pueden ocultarse por perfil o por el espacio de trabajo.", "Si una vista no aparece, confirme el perfil y sus permisos en Configuración.", "La navegación no genera documentos; las búsquedas no modifican datos.", "Permite entrar al procedimiento operativo correspondiente.", [("21-configuracion.png", "Centro de configuración y secciones disponibles"), ("32-perfiles-permisos.png", "Perfiles y matriz de permisos efectivos")], fig)
    fig = add_procedure(doc, "3.2 Usar el área de pruebas", "Practicar la cadena documental sin modificar información operativa.", "Administrador, Auditor, Supervisor y Vendedor cuando el permiso sandbox está habilitado.", "El administrador debe mantener habilitada el área de pruebas.", "Botón ? o selector Espacio de trabajo > Simulación guiada.", ["Abra la Simulación guiada.", "Pulse Ejecutar paso para generar Cotización, OC cliente, OC proveedor, Entrada, Salida y Remisión.", "Revise que todos los folios comiencen con TEST.", "Consulte el historial o previsualice el expediente.", "Use Restablecer para borrar únicamente esta simulación."], "El contador llega a 6 de 6 y la operación real conserva sus cantidades.", "El sistema separa el almacén simulado del almacén operativo.", "Si la opción está oculta, el perfil no tiene permiso sandbox o el administrador la deshabilitó.", "Folios y comprobantes TEST, usuario, fecha y secuencia de eventos.", "Facilita capacitación y pruebas sin validez comercial.", [("22-area-pruebas.png", "Simulación guiada antes de ejecutar el flujo"), ("23-area-pruebas-completada.png", "Simulación completada con historial TEST")], fig)

    doc.add_heading("4 Cotizaciones", level=1)
    fig = add_procedure(doc, "4.1 Crear y emitir una cotización", "Crear una propuesta con cliente, condiciones, partidas, cantidades y fechas de entrega.", "Vendedor, Supervisor o Administrador con quotation_manage.", "Cliente identificable y productos disponibles en el catálogo comercial o captura manual permitida.", "Cotizaciones > Nueva cotización.", ["Busque un cliente frecuente o capture cliente, población, contacto, teléfono de 10 dígitos y correo.", "Defina vigencia, fecha de pago, condiciones, política de precios y cancelación.", "Busque un producto por catálogo, SKU, descripción, marca o proveedor.", "Agregue cada partida y confirme cantidad, entrega, precio y descuento.", "Revise el borrador y los totales por moneda.", "Pulse Emitir cotización y corrija cualquier campo marcado en rojo."], "Se genera un folio COT, la cotización aparece en el historial y queda disponible para PDF, Excel o conversión.", "Se requiere cliente, identidad de contacto, teléfono válido, al menos una partida, cantidad positiva y precio válido.", "Si el botón permanece deshabilitado, agregue una partida. Si el teléfono falla, capture exactamente 10 dígitos.", "Cotización, condiciones, totales, tipo de cambio protegido cuando aplica y huella de auditoría.", "La cotización vigente puede convertirse en OC del cliente.", [("04-nueva-cotizacion.png", "Formulario y borrador de una nueva cotización"), ("02-resumen-cotizaciones.png", "Resumen de cotizaciones vigentes convertidas y vencidas")], fig)
    fig = add_procedure(doc, "4.2 Consultar descargar y convertir", "Revisar la propuesta emitida y generar el pedido de venta.", "Consulta: perfiles con quotation_view. Conversión: Vendedor, Supervisor o Administrador con client_order_manage.", "Cotización vigente, cliente completo y partidas válidas.", "Cotizaciones > Panel de cotizaciones.", ["Busque por folio, cliente, contacto o producto.", "Filtre por vigentes, convertidas o vencidas.", "Use las acciones para previsualizar PDF o exportar Excel.", "Pulse Convertir y revise la lista de bloqueos y advertencias.", "Confirme la conversión cuando todos los bloqueos estén resueltos.", "En OC de clientes, registre la aceptación para activar el surtido."], "Se crea una OC de cliente vinculada al folio y la huella de la cotización.", "Bloquean la conversión: permiso insuficiente, cancelación, conversión previa, vigencia vencida, sin partidas, partidas inválidas o sin cliente. La falta de contacto o aceptación se muestra como advertencia.", "Una cotización vencida no se edita como versión; emita un nuevo folio vigente. El sistema no importa de regreso el Excel de cotización en esta versión.", "PDF interno, Excel espejo, referencia de conversión y OC de cliente.", "La OC queda pendiente de aceptación y activación.", [("03-panel-cotizaciones.png", "Historial con búsqueda filtros y acciones de cotización"), ("05-oc-clientes.png", "Órdenes de compra de clientes derivadas de cotizaciones")], fig)

    doc.add_heading("5 Pedidos requisiciones y compras", level=1)
    fig = add_procedure(doc, "5.1 Activar la OC del cliente", "Validar la aceptación comercial antes de liberar compras o reservas.", "Supervisor o Administrador con client_order_activate. El Vendedor puede crear y consultar la OC, pero no activar el surtido si se retiró ese permiso.", "OC de cliente derivada de una cotización.", "OC de clientes > acción Activar.", ["Abra la OC pendiente.", "Seleccione OC del cliente, cotización firmada, correo de autorización o anticipo.", "Capture una referencia o adjunte la evidencia.", "Revise condiciones, partidas y totales.", "Pulse Validar y activar."], "La OC cambia a Activa y el sistema crea o actualiza la requisición con reservas y necesidades de compra.", "No se activa sin referencia verificable o archivo. La cotización ya convertida no genera una segunda OC.", "Si falta evidencia, capture el folio del documento o adjunte un archivo permitido.", "Método de aceptación, referencia, archivo, huella, usuario y fecha.", "Habilita la consolidación de compra y el seguimiento operativo.", [("05-oc-clientes.png", "Panel de OC de clientes y estado de activación")], fig)
    fig = add_procedure(doc, "5.2 Consolidar y confirmar OC a proveedor", "Agrupar faltantes compatibles sin perder la relación con cada pedido de cliente.", "Supervisor o Administrador con supplier_order_manage.", "OC de clientes activa, producto relacionado con proveedor, moneda, almacén, fecha requerida y condiciones compatibles.", "OC a proveedores.", ["Revise las propuestas preparadas automáticamente.", "Atienda el bloque Por relacionar cuando falte producto o proveedor.", "Abra el detalle y confirme las asignaciones por pedido.", "Verifique proveedor, moneda, almacén, fecha y términos de pago.", "Confirme la OC; la versión queda bloqueada y auditable.", "Descargue o previsualice el documento."], "Se genera una OC a proveedor con una o varias asignaciones y cada partida conserva su requisición de origen.", "No se compra para un pedido de cliente cuya OC no esté activa. Las partidas sin proveedor quedan separadas.", "Si una partida aparece Por relacionar, complete producto y proveedor antes de confirmar.", "OC proveedor, versión, asignaciones, aprobación, usuario, fecha y conciliación de tres vías.", "Habilita la recepción contra la cantidad comprada.", [("06-oc-proveedores.png", "Propuestas y órdenes confirmadas por proveedor"), ("26-expediente-compras.png", "Compras relacionadas dentro del expediente")], fig)

    doc.add_heading("6 Recepciones inventario y faltantes", level=1)
    fig = add_procedure(doc, "6.1 Registrar recepción completa o parcial", "Registrar material recibido contra una compra formal y generar la entrada de almacén.", "Supervisor o Administrador con receive.", "Cantidad comprada pendiente, remisión o documento de llegada y evidencia.", "Seguimiento integral > abrir expediente > Resumen > Registrar recepción.", ["Capture la cantidad recibida en este avance.", "Confirme proveedor y remisión; si deja el folio vacío, el sistema toma el nombre del archivo adjunto.", "Seleccione almacén, cantidad rechazada y estado de calidad.", "Capture lote, serie y caducidad cuando apliquen.", "Adjunte la remisión o evidencia obligatoria.", "Si la cantidad aceptada es menor al pendiente, complete Tratamiento del faltante.", "Registre la recepción."], "Se crea REC, se actualiza la cantidad recibida y se genera la entrada de inventario por la cantidad aceptada.", "No se admite cantidad cero, excedente sin permiso especial, evidencia ausente ni recepción parcial sin decisión del faltante.", "Si el material está dañado, capture la cantidad rechazada y el estado Rechazada o Cuarentena; no lo trate como disponible.", "Remisión del proveedor, evidencia con huella, lote, serie, caducidad, calidad, usuario y fecha.", "Actualiza disponibilidad y puede habilitar reserva, surtido y remisión.", [("30-recepcion-parcial-faltante.png", "Recepción formal con campos de calidad y tratamiento del faltante"), ("27-expediente-recepciones.png", "Recepciones y factura del proveedor en el expediente")], fig)
    doc.add_heading("6.2 Resolver un faltante", level=2)
    doc.add_paragraph("Cuando la cantidad aceptada es menor al pendiente, el sistema exige seleccionar una de estas alternativas antes de guardar:")
    shortages = [
        ("same_order", "Llegará posteriormente en la misma OC"),
        ("new_shipment", "Llegará en otro embarque"),
        ("new_supplier_order", "Se relacionará con otra OC"),
        ("other_supplier", "Se comprará con otro proveedor"),
        ("substitution", "Se sustituirá con autorización"),
        ("supplier_unavailable", "El proveedor no puede entregar"),
        ("authorized_cancellation", "Cancelar saldo autorizado"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.cell(0, 0).text = "Clave"
    table.cell(0, 1).text = "Decisión operativa"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
    for i, row in enumerate(shortages):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = row
        if i % 2:
            set_cell_shading(cells[0], "F3F8FA")
            set_cell_shading(cells[1], "F3F8FA")
        for cell in cells:
            set_cell_margins(cell)
    set_table_borders(table)
    doc.add_paragraph("Además de la decisión, registre fecha compromiso y notas con responsable, motivo, referencia o autorización. El expediente mostrará el pedido afectado, la alerta, la próxima acción y la línea de tiempo.")
    fig = add_procedure(doc, "6.3 Consultar existencias reservas movimientos y conteos", "Controlar disponibilidad por almacén y conservar movimientos auditables.", "Consulta: perfiles con inventory_view. Gestión: Supervisor o Administrador con inventory_manage.", "Producto de inventario y almacén identificado.", "Almacén > Existencias, Movimientos o Conteos físicos.", ["En Existencias, busque por producto o SKU y filtre almacén, estado y moneda.", "En Movimientos, filtre entradas, salidas, traspasos o mermas.", "Para una entrada manual, seleccione un producto relacionado con proveedor y adjunte cotización.", "Para salida, traspaso o merma, confirme existencia suficiente.", "En Conteos físicos, capture sólo referencias con existencia y aplique ajustes con autorización."], "Las existencias y el kardex reflejan movimientos formales y ajustes permitidos.", "El sistema impide salidas mayores a la existencia y exige respaldo de compra para una entrada manual.", "Si no aparece un producto, concilie primero su catálogo y proveedor. Si no hay stock, complete la compra y recepción.", "Movimiento, almacén, cantidad, documento origen, evidencia y usuario.", "Afecta reservas, surtido y posibilidad de generar remisión.", [("07-resumen-almacen.png", "Resumen de almacén y operaciones disponibles"), ("08-inventario.png", "Existencias por almacén estado y moneda"), ("09-movimientos.png", "Kardex de entradas salidas traspasos y mermas"), ("10-conteos.png", "Conteo físico de referencias con existencia")], fig)

    doc.add_heading("7 Seguimiento por pedido y por partida", level=1)
    fig = add_procedure(doc, "7.1 Seguimiento por pedido general", "Revisar el avance agregado de todas las partidas de una requisición.", "Administrador y Supervisor; otros perfiles según permisos de consulta.", "Pedido o requisición creado.", "Seguimiento y remisiones > Seguimiento integral.", ["Busque por pedido, cliente, referencia o producto.", "Filtre estado, etapa, periodo, responsable y vencimiento.", "Revise solicitadas, compradas, recibidas y entregadas.", "Abra el expediente para consultar Comercial, Compras, Recepciones, Almacén, Entregas, Documentos y Línea de tiempo."], "El usuario identifica el cuello de botella global y la próxima acción del pedido.", "Las fechas vencidas y partidas bloqueadas generan alertas en el expediente.", "Si el avance global parece incompleto, abra cada partida; una sola partida pendiente mantiene abierto el pedido.", "Resumen, documentos relacionados, responsables, compromisos y auditoría.", "Dirige al seguimiento detallado o a la operación formal necesaria.", [("12-seguimiento-pedidos.png", "Lista de seguimiento general con filtros y avance"), ("24-expediente-resumen-partidas.png", "Expediente con indicadores globales y cantidades por partida")], fig)
    fig = add_procedure(doc, "7.2 Seguimiento por partida", "Localizar el estado exacto de un producto dentro del pedido.", "Administrador y Supervisor; Vendedor en sus operaciones cuando el permiso lo permita.", "Expediente abierto.", "Seguimiento integral > Abrir expediente > Resumen.", ["Ubique la partida por catálogo y descripción.", "Compare Solicitadas, Reservadas, Compradas, Recibidas, Entregadas y Pendientes.", "Abra las pestañas Compras, Recepciones, Almacén y Entregas para rastrear documentos.", "Use Línea de tiempo para revisar responsable, fecha, referencia, próxima acción y bloqueo.", "Use Registrar avance sólo para seguimiento manual; use Registrar recepción o Generar remisión para movimientos formales."], "Cada cantidad puede rastrearse hasta la OC, recepción, almacén, reserva, salida o remisión correspondiente.", "No confunda una nota de seguimiento con una recepción o salida. Sólo las operaciones formales cambian cantidades.", "Si una partida no tiene proveedor o documento, atienda el aviso y complete la etapa anterior.", "Referencias cruzadas, cantidades, lotes, evidencias y línea de tiempo.", "Permite resolver la partida sin perder el contexto del pedido general.", [("25-expediente-comercial.png", "Origen comercial y pedido de venta"), ("26-expediente-compras.png", "OC y asignación de compra"), ("27-expediente-recepciones.png", "Recepción física y factura del proveedor"), ("28-expediente-linea-tiempo.png", "Línea de tiempo auditable")], fig)

    doc.add_heading("8 Registrar avance", level=1)
    fig = add_procedure(doc, "8.1 Registrar seguimiento manual", "Documentar contacto, cambio de estado, compromiso, documento o incidencia sin alterar inventario.", "Perfiles con permiso edit. La reasignación a otra persona requiere permiso de compra, normalmente Supervisor o Administrador.", "Partida existente y expediente abierto.", "Seguimiento integral > Abrir expediente > Registrar avance.", ["Revise el contexto automático: pedido, partida, estado, responsable, compromiso y cantidades.", "Seleccione Seguimiento realizado, Cambio de estado, Compromiso o próxima acción, Documento recibido, Incidencia o bloqueo, u Otro.", "Capture una descripción breve obligatoria.", "Para compromiso, capture próxima acción y fecha.", "Para documento o incidencia, adjunte evidencia; la incidencia también exige bloqueo y prioridad.", "Confirme próximo responsable y guarde."], "La bitácora agrega un evento manual con usuario, fecha, resultado, responsable y compromiso. Las cantidades no cambian.", "El sistema exige evidencia para Documento e Incidencia, y restringe la reasignación según permiso.", "Si necesita registrar llegada o entrega física, cierre el modal y use la operación formal correspondiente.", "Nota, evidencia, huella, estado de seguimiento, responsable y fecha compromiso.", "Actualiza alertas y próxima acción sin sustituir OC, recepción, salida o remisión.", [("29-registrar-avance.png", "Modal Registrar avance con campos dinámicos y separación de operaciones formales")], fig)

    doc.add_heading("9 Remisiones salida y entrega al cliente", level=1)
    fig = add_procedure(doc, "9.1 Generar remisión y salida", "Entregar una cantidad disponible sin duplicar la salida ni exceder lo pendiente.", "Supervisor o Administrador con deliver.", "Pedido activo, cantidad pendiente y producto disponible por recepción o reserva.", "Seguimiento integral > Abrir expediente > Resumen > Generar remisión.", ["Revise el pendiente formal y capture la cantidad a entregar.", "Confirme el folio generado automáticamente por PROBIOLAB.", "Capture nombre de quien recibe y fecha programada.", "Adjunte remisión firmada, fotografía o firma obligatoria.", "Registre observaciones si existe entrega parcial, rechazo o incidencia.", "Pulse Generar remisión."], "Se crea el folio REM, la salida de almacén y la entrega parcial o total; el expediente vuelve a abrir con confirmación.", "No se entrega más de lo disponible o pendiente salvo autorización especial registrada. La evidencia y el receptor son obligatorios.", "Si no aparece Generar remisión, complete recepción o reserva. Si falta inventario, atienda compra y recepción.", "Remisión, salida, receptor, fecha, evidencia y huella.", "Reduce pendiente de entrega y puede habilitar facturación o cierre.", [("31-generar-remision.png", "Generación formal de remisión con receptor y evidencia"), ("13-entregas-remisiones.png", "Panel de entregas pendientes parciales y completas")], fig)
    doc.add_heading("9.2 Estados disponibles y límites de la versión", level=2)
    doc.add_paragraph("La versión validada presenta pedidos Pendientes, Parciales o Completos en el panel de entregas y remisiones; cada remisión se registra como entrega parcial o total. No existen todavía pantallas independientes para Borrador, Preparando, Autorizada, En tránsito, Rechazada o Cancelada. Tampoco hay un módulo formal separado de devoluciones y reposiciones. Estos puntos figuran como pendientes en el reporte de pruebas; no deben simularse con notas genéricas.")

    doc.add_heading("10 Expediente auditoría catálogos y administración", level=1)
    fig = add_procedure(doc, "10.1 Consultar expediente y auditoría", "Reconstruir la cadena completa de una operación.", "Expediente: perfiles autorizados. Auditoría consolidada: Administrador y Auditor.", "Folio, cliente, pedido o referencia.", "Seguimiento integral > Expediente; o Auditoría.", ["Abra el expediente del pedido.", "Recorra Comercial, Compras, Recepciones, Almacén, Entregas, Documentos y Línea de tiempo.", "En Auditoría, busque por folio, documento, movimiento, usuario o detalle.", "Filtre categoría y fechas.", "Previsualice o descargue los documentos disponibles."], "La consulta muestra origen, versiones disponibles, cambios de responsables, fechas, cancelaciones y documentos relacionados.", "Auditoría se oculta al Vendedor y en el espacio de pruebas.", "Si un documento no aparece, revise la etapa que debía generarlo y el permiso del perfil.", "Cotizaciones, OC, recepciones, facturas, movimientos, remisiones y bitácoras.", "Sustenta supervisión, aclaraciones y cierre.", [("20-auditoria.png", "Auditoría consolidada con filtros"), ("28-expediente-linea-tiempo.png", "Eventos responsables y referencias del expediente")], fig)
    fig = add_procedure(doc, "10.2 Administrar productos proveedores listas clientes y reportes", "Mantener los datos maestros que alimentan la operación.", "Vendedor gestiona sus clientes y cotizaciones. Supervisor y Administrador gestionan catálogos según permisos.", "Permiso correspondiente y justificación cuando el cambio sea sensible.", "Productos, Proveedores, Listas de precios, Clientes, Reportes y Configuración.", ["En Productos, consulte catálogo, fotografías, etiquetas, proveedor y existencias.", "En Proveedores, cree o edite datos y justifique eliminaciones.", "En Listas de precios, seleccione proveedor, cargue Excel o CSV, revise conflictos y aplique la lista.", "En Clientes, consulte expediente, crédito, bloqueo y saldo a favor.", "En Reportes, abra inventario, mermas o evidencia.", "En Configuración, administre usuarios, permisos, divisas y tema."], "Los módulos operativos utilizan catálogos conciliados y permisos efectivos.", "Las listas no crean inventario; sólo actualizan el catálogo comercial. El correo debe ser válido y siempre debe quedar un administrador activo.", "Si una lista no coincide, corrija proveedor, moneda o descripción antes de aplicarla.", "Historial de listas, cambios de proveedor y cliente, permisos y auditoría.", "Reduce duplicados y evita compras o cotizaciones con datos incompletos.", [("14-resumen-productos.png", "Resumen de productos y accesos administrativos"), ("15-catalogo-productos.png", "Catálogo de productos con proveedor precio y existencia"), ("16-proveedores.png", "Proveedores y contactos"), ("17-listas-precios.png", "Carga y conciliación de listas de precios"), ("18-clientes.png", "Cartera de clientes y expediente"), ("19-reportes.png", "Reportes operativos disponibles")], fig)

    doc.add_heading("11 Problemas frecuentes", level=1)
    problems = [
        ("No puedo convertir una cotización", "Permiso, vigencia, cancelación, conversión previa, cliente o partidas", "Abra el detalle del bloqueo; solicite permiso, emita un nuevo folio vigente o complete cliente y partidas."),
        ("Falta evidencia", "El flujo formal exige archivo o referencia verificable", "Adjunte PDF o imagen permitida. En aceptación puede usar referencia; en recepción y entrega el archivo es obligatorio."),
        ("Una partida no tiene proveedor", "El producto no está relacionado con un proveedor", "Abra Productos o Listas de precios, concilie el artículo y vuelva a OC a proveedores."),
        ("El proveedor entregó incompleto", "Cantidad aceptada menor al pendiente", "Registre recepción parcial y seleccione el tratamiento del faltante, fecha y responsable."),
        ("El pedido aparece bloqueado", "Incidencia, fecha vencida, falta de aceptación o proveedor", "Abra el expediente y la Línea de tiempo; atienda la alerta específica."),
        ("No puedo generar una remisión", "Sin disponibilidad, sin pendiente o permiso insuficiente", "Complete recepción o reserva, confirme el pendiente y solicite permiso deliver."),
        ("No hay inventario disponible", "Existencia libre insuficiente", "Consolide compra, registre recepción liberada y vuelva al expediente."),
        ("La entrega fue parcial", "Queda cantidad pendiente", "Conserve la evidencia de la primera remisión y programe una nueva remisión sólo por el saldo."),
        ("El usuario no tiene permiso", "El perfil o una excepción individual niega la capacidad", "El Administrador revisa Configuración > Perfiles y permisos."),
        ("La fecha compromiso está vencida", "La partida conserva una fecha anterior a hoy", "Registre un compromiso realista, próximo responsable y causa del retraso."),
    ]
    table = doc.add_table(rows=1, cols=3)
    headers = ["Caso", "Causa probable", "Solución"]
    for j, header in enumerate(headers):
        table.cell(0, j).text = header
        set_cell_shading(table.cell(0, j), NAVY)
        for run in table.cell(0, j).paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(problems):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = value
            set_cell_margins(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2:
                set_cell_shading(cells[j], "F3F8FA")
    set_repeat_table_header(table)
    set_table_borders(table)

    doc.add_heading("12 Recorrido para un usuario nuevo", level=1)
    add_numbered_steps(doc, [
        "Abra el área de pruebas desde el botón ? y complete los seis pasos TEST.",
        "En Productos y Proveedores, confirme que el artículo tenga proveedor y precio.",
        "Cree una cotización con cliente, teléfono, condiciones y una partida.",
        "Emita la cotización, previsualice el PDF y exporte el Excel.",
        "Convierta la cotización y registre aceptación u OC del cliente.",
        "Active la OC del cliente y revise la requisición creada.",
        "Confirme la propuesta de OC al proveedor y sus asignaciones.",
        "Registre recepción; si es parcial, documente la decisión del faltante.",
        "Revise entrada, existencias y reserva en el expediente.",
        "Use Registrar avance para una nota de seguimiento sin cambiar cantidades.",
        "Genere la remisión sólo por la cantidad disponible; adjunte evidencia de entrega.",
        "Revise la entrega, la factura y la Línea de tiempo antes de cerrar el expediente.",
    ])

    doc.add_heading("13 Glosario de términos y folios", level=1)
    glossary = [
        ("COT", "Cotización comercial emitida por PROBIOLAB."),
        ("OC cliente", "Orden o aceptación que autoriza el pedido de venta."),
        ("OCP", "Orden de compra a proveedor."),
        ("REC", "Recepción física contra una compra."),
        ("INV IN", "Entrada formal de almacén."),
        ("INV OUT", "Salida formal vinculada a remisión."),
        ("REM", "Remisión emitida para entrega al cliente."),
        ("FAC", "Factura al cliente."),
        ("Partida", "Producto y cantidad específicos dentro de un pedido."),
        ("Reserva", "Cantidad comprometida para un pedido sin haber salido aún."),
        ("Huella", "Resumen de integridad utilizado para identificar evidencia o documento."),
        ("TEST", "Prefijo exclusivo del entorno de entrenamiento sin validez operativa."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Término"
    table.cell(0, 1).text = "Definición"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for i, (term, definition) in enumerate(glossary):
        cells = table.add_row().cells
        cells[0].text = term
        cells[1].text = definition
        cells[0].paragraphs[0].runs[0].bold = True
        for cell in cells:
            set_cell_margins(cell)
            if i % 2:
                set_cell_shading(cell, "F3F8FA")
    set_table_borders(table)

    doc.add_heading("14 Control de cambios", level=1)
    table = doc.add_table(rows=1, cols=4)
    for j, text in enumerate(["Versión", "Fecha", "Descripción", "Responsable"]):
        table.cell(0, j).text = text
        set_cell_shading(table.cell(0, j), NAVY)
        for run in table.cell(0, j).paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    row = table.add_row().cells
    values = ["1.0", TODAY, "Emisión inicial validada contra la versión funcional; incluye capturas, procedimientos, límites conocidos y correcciones de QA.", "Equipo PROBIOLAB"]
    for i, value in enumerate(values):
        row[i].text = value
        set_cell_margins(row[i])
    set_table_borders(table)

    add_header_footer(doc, "Manual de Usuario Bio Probiolab")
    OUT.mkdir(parents=True, exist_ok=True)
    prevent_table_row_splits(doc)
    doc.save(MANUAL)
    return MANUAL


def build_report():
    doc = Document()
    configure_styles(doc)
    add_cover(doc, "Reporte de Pruebas y Correcciones Bio Probiolab", "Validación funcional visual y documental de la versión 1.0")
    add_toc(doc)
    doc.add_heading("1 Resumen ejecutivo", level=1)
    doc.add_paragraph("La validación cerró con 78 pruebas automatizadas del repositorio y 65 comprobaciones de interfaz aprobadas. Se generaron 32 capturas a 1440 por 1000 píxeles en un perfil aislado. Se corrigieron cuatro defectos reproducibles. No se alteró información productiva.")
    doc.add_paragraph("La validación del servidor central PostgreSQL quedó limitada a las pruebas automatizadas con base temporal porque el equipo no tiene DATABASE_URL, PostgreSQL ni Docker configurados. El modo local sí fue recorrido de extremo a extremo con datos sintéticos.")

    doc.add_heading("2 Arquitectura e inventario inspeccionado", level=1)
    inventory = [
        ("Interfaz", "index.html, app.js, orders.js, estilos y componentes modales"),
        ("Reglas operativas", "orders-core.js, quotation-conversion.js y form-validation.js"),
        ("Acceso", "access.js con Administrador, Auditor, Supervisor y Vendedor"),
        ("Persistencia", "persistence.js y API REST en server/server.js"),
        ("Base central", "PostgreSQL configurado mediante DATABASE_URL"),
        ("Documentos", "flow-documents.js, vista previa y descargas PDF o Excel"),
        ("Pruebas", "78 casos Node y 65 comprobaciones de navegador"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Componente"
    table.cell(0, 1).text = "Cobertura"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(inventory):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = row
        for cell in cells:
            set_cell_margins(cell)
            if i % 2:
                set_cell_shading(cell, "F3F8FA")
    set_table_borders(table)

    fixes_heading = doc.add_heading("3 Errores corregidos", level=1)
    fixes_heading.paragraph_format.page_break_before = True
    fixes = [
        ("COR-01", "Datos de demostración", "El expediente marcado como completo no conservaba aceptación ni evidencia sintética de recepción y entrega.", "Agregar aceptación con referencia y huella; recepción con cantidades aceptadas, lote y evidencia; remisión con receptor, estado y evidencia.", "Pruebas demo-flow y expediente visual", "Corregido"),
        ("COR-02", "Área de pruebas", "La guía decía Recepción parcial aunque el paso recibía 5 de 5 unidades.", "Cambiar la etiqueta a Recepción y entrada y describir el movimiento simulado.", "Regresión textual dedicada", "Corregido"),
        ("COR-03", "Navegación", "La pantalla Área de pruebas existía pero no tenía ruta visible; el botón ? sólo abría el selector.", "Agregar Simulación guiada al selector y conectar el botón ? con la vista, respetando permiso sandbox.", "Prueba de navegación y comprobación UI HELP-SANDBOX", "Corregido"),
        ("COR-04", "Dependencia de iconos", "Sin acceso al CDN, lucide no existía y una llamada no protegida generaba error de ejecución.", "Instalar un adaptador vacío cuando la biblioteca remota no carga.", "Prueba unitaria y consola de navegador sin errores", "Corregido"),
    ]
    table = doc.add_table(rows=1, cols=6)
    headers = ["ID", "Módulo", "Resultado actual", "Corrección", "Prueba", "Estado"]
    for j, header in enumerate(headers):
        table.cell(0, j).text = header
        set_cell_shading(table.cell(0, j), NAVY)
        for run in table.cell(0, j).paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(fixes):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = value
            set_cell_margins(cells[j], 80, 75, 80, 75)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2:
                set_cell_shading(cells[j], "F3F8FA")
    set_repeat_table_header(table)
    set_table_borders(table)

    doc.add_heading("4 Pendientes y límites conocidos", level=1)
    pending = [
        ("PEN-01", "Infraestructura", "Conectar y probar una instancia PostgreSQL real mediante DATABASE_URL.", "Alta", "Operación central y multiusuario"),
        ("PEN-02", "Remisiones", "Implementar estados Borrador, Preparando, Autorizada, En tránsito, Rechazada y Cancelada con transiciones y autorizaciones.", "Alta", "Trazabilidad logística avanzada"),
        ("PEN-03", "Devoluciones", "Crear flujo formal para rechazo del cliente, devolución, reposición y reingreso de inventario.", "Alta", "Posventa e inventario"),
        ("PEN-04", "Roles", "Separar perfiles de Compras y Almacén; actualmente esas responsabilidades se cubren con Supervisor o permisos personalizados.", "Media", "Segregación de funciones"),
        ("PEN-05", "Cotizaciones", "Importar y conciliar el Excel devuelto por el cliente y administrar revisiones bajo un folio base.", "Media", "Versiones y aceptación"),
        ("PEN-06", "Operación sin Internet", "Empaquetar localmente jsPDF, XLSX, PDF.js e iconos; la interfaz tolera la falta de iconos, pero algunas exportaciones dependen de CDN.", "Media", "Continuidad offline"),
        ("PEN-07", "Lotes", "Extender el kardex para consultar existencias por lote, serie y caducidad, no sólo conservarlos en la recepción.", "Media", "Trazabilidad física"),
    ]
    table = doc.add_table(rows=1, cols=5)
    for j, header in enumerate(["ID", "Área", "Pendiente", "Prioridad", "Impacto"]):
        table.cell(0, j).text = header
        set_cell_shading(table.cell(0, j), NAVY)
        for run in table.cell(0, j).paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(pending):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = value
            set_cell_margins(cells[j])
            if i % 2:
                set_cell_shading(cells[j], "F3F8FA")
    set_repeat_table_header(table)
    set_table_borders(table)

    doc.add_heading("5 Matriz de pruebas", level=1)
    matrix = [
        ("PR-01", "Cotización completa y conversión", "Automatizada y UI", "Aprobada", "Folio, partidas y OC vinculada"),
        ("PR-02", "Cotización bloqueada por datos", "Automatizada", "Aprobada", "Mensajes específicos por bloqueo"),
        ("PR-03", "Bloqueo por permisos vigencia o conversión previa", "Automatizada", "Aprobada", "Códigos permission_denied, expired y already_converted"),
        ("PR-04", "Compra consolidada para varios pedidos", "Automatizada", "Aprobada", "Asignaciones conservan requisición de origen"),
        ("PR-05", "Recepción completa", "Automatizada y demo", "Aprobada", "REC y entrada vinculadas"),
        ("PR-06", "Recepción parcial", "UI de extremo a extremo", "Aprobada", "4 iniciales más 3 aceptadas igual a 7"),
        ("PR-07", "Faltante con siete alternativas", "UI y código", "Aprobada", "Siete opciones más opción inicial"),
        ("PR-08", "Reserva y salida de inventario", "Automatizada", "Aprobada", "No compromete existencias reservadas"),
        ("PR-09", "Remisión parcial", "UI de extremo a extremo", "Aprobada", "Entrega adicional de 1 unidad y evidencia"),
        ("PR-10", "Remisión completa", "Automatizada y demo", "Aprobada", "Pedido llega a entrega total"),
        ("PR-11", "Entrega con evidencia", "UI y demo", "Aprobada", "Receptor, archivo y huella"),
        ("PR-12", "Incidencia cancelación y reasignación", "Automatizada y código", "Aprobada", "Motivo, permisos y bitácora"),
        ("PR-13", "Seguimiento por pedido", "UI", "Aprobada", "Filtros, avance y expediente"),
        ("PR-14", "Seguimiento por partida", "UI", "Aprobada", "Cantidades y pestañas relacionadas"),
        ("PR-15", "Registrar avance", "UI", "Aprobada", "Modo tracking separado de formal"),
        ("PR-16", "Alertas de vencimiento", "UI", "Aprobada", "Fecha vencida visible en expediente"),
        ("PR-17", "Permisos por rol", "Automatizada y UI", "Aprobada", "Vendedor no ve Auditoría"),
        ("PR-18", "Auditoría y línea de tiempo", "Automatizada y UI", "Aprobada", "Documentos y eventos consolidados"),
        ("PR-19", "Prevención de duplicados", "Automatizada", "Aprobada", "Importación y carga idempotentes"),
        ("PR-20", "Área TEST aislada", "UI de extremo a extremo", "Aprobada", "El almacén operativo no cambia"),
        ("PR-21", "Navegación de 22 vistas", "UI", "Aprobada", "Encabezados y rutas visibles"),
        ("PR-22", "Consola del navegador", "UI", "Aprobada", "Sin errores de ejecución"),
        ("PR-23", "PostgreSQL real", "Entorno", "Pendiente", "No existe DATABASE_URL ni servicio local"),
    ]
    table = doc.add_table(rows=1, cols=5)
    for j, header in enumerate(["ID", "Escenario", "Tipo", "Resultado", "Evidencia"]):
        table.cell(0, j).text = header
        set_cell_shading(table.cell(0, j), NAVY)
        for run in table.cell(0, j).paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(matrix):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = value
            set_cell_margins(cells[j], 75, 70, 75, 70)
            if i % 2:
                set_cell_shading(cells[j], "F3F8FA")
    set_repeat_table_header(table)
    set_table_borders(table)

    doc.add_heading("6 Archivos modificados", level=1)
    files = [
        "app.js - tolerancia a la ausencia de la biblioteca de iconos",
        "demo-flow.js - aceptación y evidencias consistentes en el expediente demostrativo",
        "index.html - acceso visible a Simulación guiada y texto correcto de recepción",
        "workspace.js - visibilidad y cierre del acceso según permiso sandbox",
        "tests/demo-flow.test.js - regresiones de aceptación, evidencias y guía",
        "tests/navigation.test.js - regresión de la ruta visible de pruebas",
        "tests/theme.test.js - regresión de operación sin iconos remotos",
        "scripts/manual-qa-capture.cjs - recorrido UI aislado y capturas",
        "scripts/build-user-manual.py - generación reproducible de entregables",
    ]
    for value in files:
        doc.add_paragraph(value, style="List Bullet")

    doc.add_heading("7 Evidencias y comandos", level=1)
    add_label_paragraph(doc, "Pruebas del repositorio", "npm test: 78 aprobadas, 0 fallidas.")
    add_label_paragraph(doc, "Regresión UI", "scripts/manual-qa-capture.cjs: 65 aprobadas, 32 capturas.")
    add_label_paragraph(doc, "Resolución de captura", "1440 x 1000 píxeles; perfil Chromium aislado; datos sintéticos QA y DEMO.")
    add_label_paragraph(doc, "Resultado", "No se detectaron errores de ejecución en la consola después de las correcciones.")

    doc.add_heading("8 Control de cambios", level=1)
    table = doc.add_table(rows=1, cols=4)
    for j, text in enumerate(["Versión", "Fecha", "Cambio", "Estado"]):
        table.cell(0, j).text = text
        set_cell_shading(table.cell(0, j), NAVY)
        for run in table.cell(0, j).paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    row = table.add_row().cells
    for i, value in enumerate(["1.0", TODAY, "Validación inicial, cuatro correcciones y registro de pendientes.", "Emitido"]):
        row[i].text = value
        set_cell_margins(row[i])
    set_table_borders(table)

    add_header_footer(doc, "Reporte de Pruebas y Correcciones Bio Probiolab")
    OUT.mkdir(parents=True, exist_ok=True)
    prevent_table_row_splits(doc)
    doc.save(REPORT)
    return REPORT


if __name__ == "__main__":
    print(build_manual())
    print(build_report())
